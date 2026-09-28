"""Bounded, isolated parsing and malware scanning; no OCR or active content execution."""
import hashlib
import io
import json
import re
import socket
import struct
import subprocess
import sys
import zipfile
from pathlib import Path
from ..errors import fail

FORMATS = {'.txt': 'text/plain', '.md': 'text/markdown', '.pdf': 'application/pdf',
           '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}
BLOCKED_NOTICES = [
    re.compile(r'\bDeloitte\s+Accounting\s+Research\s+Tool\b', re.I),
    re.compile(r'dart\.deloitte\.com', re.I),
    re.compile(r'\bFinancial\s+Accounting\s+Foundation\b.{0,100}\bAll\s+rights\s+reserved\b', re.I | re.S),
]


def safe_filename(name):
    name = name.replace('\\', '/').split('/')[-1]
    name = re.sub(r'[\x00-\x1f\x7f]', '', name).strip()
    return (name or 'document.txt')[:180]


def preflight(data, filename, config):
    ext = Path(filename).suffix.lower()
    if ext not in FORMATS:
        fail('FILE_TYPE', 'Use text, Markdown, text-based PDF, or DOCX.', 415)
    if not data or len(data) > config.max_upload_mb * 1024 * 1024:
        fail('FILE_SIZE', f'Upload a non-empty file up to {config.max_upload_mb} MB.', 413)
    if ext == '.pdf' and not data.startswith(b'%PDF-'):
        fail('FILE_SIGNATURE', 'The file does not have a PDF signature.', 415)
    if ext == '.docx':
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                files = archive.infolist()
                if len(files) > 2000 or sum(f.file_size for f in files) > 40 * 1024 * 1024:
                    fail('ARCHIVE_LIMIT', 'The DOCX expands beyond the supported size.', 413)
                if 'word/document.xml' not in archive.namelist():
                    fail('FILE_SIGNATURE', 'This is not a supported DOCX file.', 415)
                if any('vbaProject' in x.filename or '..' in Path(x.filename).parts for x in files):
                    fail('ACTIVE_CONTENT', 'Macros or unsafe archive members are not supported.', 415)
                if any(f.file_size > 1000 * max(1, f.compress_size) for f in files):
                    fail('ARCHIVE_LIMIT', 'Excessive compression ratio.', 413)
        except zipfile.BadZipFile:
            fail('FILE_SIGNATURE', 'The DOCX archive is invalid.', 415)
    return FORMATS[ext]


def scan(data, config):
    if not config.clamav_host:
        if config.upload_scanning_required:
            fail('SCANNER_UNAVAILABLE', 'Document scanning is unavailable; the upload was not accepted.', 503)
        return 'not_scanned_local_only'
    try:
        with socket.create_connection((config.clamav_host, config.clamav_port), timeout=20) as sock:
            sock.sendall(b'zINSTREAM\0')
            for offset in range(0, len(data), 65536):
                chunk = data[offset:offset + 65536]
                sock.sendall(struct.pack('!I', len(chunk)) + chunk)
            sock.sendall(struct.pack('!I', 0))
            response = sock.recv(4096).decode(errors='replace')
            if not response.rstrip('\0\n').endswith(' OK'):
                fail('UNSAFE_FILE', 'The file did not pass the malware scan.', 422)
    except OSError:
        fail('SCANNER_UNAVAILABLE', 'Document scanning is unavailable.', 503)
    return 'clean'


def parse_bytes(data, ext, character_limit):
    """Invoked in a subprocess by parse_isolated, or directly by small unit tests."""
    chunks = []
    if ext in {'.txt', '.md'}:
        text = data.decode('utf-8-sig', errors='strict')
        if '\0' in text:
            raise ValueError('Binary content in text file.')
        chunks = [{'locator': f'Paragraph {i+1}', 'text': p.strip()} for i, p in enumerate(re.split(r'\n\s*\n', text)) if p.strip()]
    elif ext == '.pdf':
        from ..pdf_parser import parse
        chunks = [{'locator': f'Page {p["pdf"]["physical_page"]}', **p}
                  for p in parse(data, page_limit=150, character_limit=character_limit)]
    elif ext == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip():
                chunks.append({'locator': f'Paragraph {i+1}', 'text': p.text})
        for i, table in enumerate(doc.tables):
            for j, row in enumerate(table.rows):
                chunks.append({'locator': f'Table {i+1}, row {j+1}', 'text': ' | '.join(c.text for c in row.cells)})
    else:
        raise ValueError('Unsupported format.')
    total = sum(len(x['text']) for x in chunks)
    if total > character_limit:
        raise ValueError('Extracted text exceeds the configured limit; split the document.')
    if total < 10:
        raise ValueError('No usable text was extracted. Scanned PDFs require a text-based replacement; OCR is not enabled.')
    # Bounded chunks preserve locator. Do not silently discard the end of the file.
    output = []
    for item in chunks:
        for offset in range(0, len(item['text']), 3500):
            chunk = {**item, 'locator': item['locator'] + (f', part {offset//3500+1}' if offset else ''),
                     'text': item['text'][offset:offset+3500]}
            if 'pdf' in item:
                chunk['pdf'] = {**item['pdf'], 'character_range': [offset, min(offset+3500, len(item['text']))]}
            output.append(chunk)
    return output


def parse_isolated(data, ext, limit):
    try:
        child = subprocess.run([sys.executable, '-m', 'app.parse_worker', ext, str(limit)],
                               input=data, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                               timeout=25, check=False)
    except subprocess.TimeoutExpired:
        fail('PARSE_TIMEOUT', 'Document parsing exceeded the time limit.', 422)
    if child.returncode == 2 and ext == '.pdf' and len(child.stderr) <= 4000:
        try:
            diagnostic = json.loads(child.stderr)
            fail('PDF_PARSE_BLOCKED', 'PDF extraction is incomplete or unsupported; no document was stored. Review the original pages; OCR is not enabled.',
                 422, reason=diagnostic['pdf_error'], physical_page=diagnostic['physical_page'])
        except (ValueError, KeyError, TypeError):
            pass
    if child.returncode != 0 or len(child.stdout) > 32_000_000:
        fail('PARSE_FAILED', 'Cannot safely extract this document. Use a smaller text-based file.', 422)
    try:
        result = json.loads(child.stdout)
    except (ValueError, UnicodeDecodeError):
        fail('PARSE_FAILED', 'The parser did not produce a valid result.', 422)
    text = '\n'.join(x['text'] for x in result)
    if any(pattern.search(text) for pattern in BLOCKED_NOTICES):
        fail('SOURCE_POLICY_BLOCK', 'This upload appears to include restricted publisher content; use an authorized source route.', 403)
    return result


def checksum(data):
    return hashlib.sha256(data).hexdigest()
