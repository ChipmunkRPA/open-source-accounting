from io import BytesIO
from html import escape
from sqlalchemy import select
from ..models import Memo, MemoRevision, Run, Evidence, Review, now
from .rights import run_artifact_access


def from_result(db, run):
    existing = db.scalar(select(Memo).where(Memo.run_id == run.id))
    if existing:
        return existing
    data = run.result
    lines = [f'# {data["title"]}', '', '> AI-assisted research draft — professional review required.', '',
             '## Summary', data['summary'], '']
    for section in data['sections']:
        lines += [f'## {section["heading"]}', section['body'], '']
    lines += ['## Claims and evidence', '']
    for claim in data['claims']:
        lines.append(f'- {claim["text"]} [{claim["basis"]}; evidence: {", ".join(claim["evidence_ids"]) or "none"}]')
    for table in data.get('tables', []):
        lines += ['', f'## {table["title"]}', ' | '.join(table['columns']), ' | '.join(['---'] * len(table['columns']))]
        lines += [' | '.join(row) for row in table['rows']]
    lines += ['', '## Limitations'] + [f'- {x}' for x in data['limitations']]
    lines += ['', '## Open questions'] + [f'- {x}' for x in data.get('open_questions', [])]
    lines += ['', '## Evidence actually available']
    for e in db.scalars(select(Evidence).where(Evidence.run_id == run.id)):
        lines.append(f'- {e.title} — {e.locator}; {e.access}; evidence ID {e.id}.')
    memo = Memo(workspace_id=run.workspace_id, run_id=run.id, title=data['title'],
                body='\n'.join(lines), created_by=run.user_id, revision=1)
    db.add(memo)
    db.flush()
    save_revision(db, memo, run.user_id)
    return memo


def save_revision(db, memo, user_id):
    db.add(MemoRevision(memo_id=memo.id, number=memo.revision, body=memo.body,
                        title=memo.title, user_id=user_id))


def serial(db, memo):
    reviews = db.scalars(select(Review).where(Review.memo_id == memo.id, Review.revision == memo.revision)).all()
    return {'id': memo.id, 'workspace_id': memo.workspace_id, 'run_id': memo.run_id,
            'title': memo.title, 'body': memo.body, 'revision': memo.revision, 'updated_at': memo.updated_at,
            'review_state': 'independently_reviewed' if any(r.kind == 'independent' for r in reviews)
                            else ('self_reviewed' if reviews else 'draft'),
            'reviews': [{'reviewer_id': r.reviewer_id, 'kind': r.kind, 'note': r.note,
                         'revision': r.revision, 'created_at': r.created_at} for r in reviews]}


def verify_access(db, memo, action='quote'):
    if memo.run_id:
        run_artifact_access(db, db.get(Run, memo.run_id), action)


def export_bytes(memo, format):
    """Programmatic exports; no remote HTML or user-supplied executable markup."""
    footer = '\n\nAI-assisted research draft. Verify authority and obtain professional review before reliance.\n'
    text = memo.body + footer
    if format == 'md':
        return text.encode('utf-8'), 'text/markdown; charset=utf-8'
    if format == 'html':
        return ('<!doctype html><html><meta charset="utf-8"><title>' + escape(memo.title) +
                '</title><body><pre style="white-space:pre-wrap;font:16px/1.6 system-ui;max-width:850px;margin:48px auto">' +
                escape(text) + '</pre></body></html>').encode(), 'text/html; charset=utf-8'
    if format == 'docx':
        from docx import Document
        from docx.shared import Inches, Pt
        document = Document()
        document.sections[0].top_margin = Inches(.8)
        document.sections[0].bottom_margin = Inches(.8)
        normal = document.styles['Normal']
        normal.font.name = 'Calibri'
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(7)
        for line in text.splitlines():
            if line.startswith('# '):
                document.add_heading(line[2:], level=0)
            elif line.startswith('## '):
                document.add_heading(line[3:], level=1)
            elif line.startswith('- '):
                document.add_paragraph(line[2:], style='List Bullet')
            else:
                document.add_paragraph(line)
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if format == 'pdf':
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.pagesizes import letter
        stream = BytesIO()
        styles = getSampleStyleSheet()
        story = []
        for line in text.splitlines():
            style = styles['BodyText']
            if line.startswith('# '):
                style, line = styles['Title'], line[2:]
            elif line.startswith('## '):
                style, line = styles['Heading2'], line[3:]
            # Built-in fonts: non-Latin font embedding remains a documented limitation.
            story.append(Paragraph(escape(line) or ' ', style))
            story.append(Spacer(1, 5))
        SimpleDocTemplate(stream, pagesize=letter, leftMargin=54, rightMargin=54,
                          topMargin=50, bottomMargin=50, title=memo.title).build(story)
        return stream.getvalue(), 'application/pdf'
    raise ValueError('Unsupported export format.')
