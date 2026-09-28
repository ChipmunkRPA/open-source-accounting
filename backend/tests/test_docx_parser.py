"""Original synthetic DOCX packages; never accepted legal revisions or source rights."""
import io
import zipfile
from types import SimpleNamespace
import pytest
from docx import Document as WordDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from sqlalchemy import select, func
from app.docx_parser import parse, DocxError, VERSION, W
from app.services.documents import parse_bytes
from app.sec_core.core import digest
from app.models import Document
from app.agents.workflows import preprocess
from conftest import activate

MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def fixture():
    doc = WordDocument()
    doc.add_paragraph('Original paragraph before table.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Year'; table.cell(0, 1).text = 'Amount'
    table.cell(1, 0).text = '2025'; table.cell(1, 1).text = '(1,250.00)'
    header = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(header)
    doc.add_paragraph('Original paragraph after table.')
    out = io.BytesIO(); doc.save(out); return out.getvalue()


def mutate(raw, change, part='word/document.xml', extra=None):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as dest:
        for name in source.namelist():
            data = source.read(name)
            if name == part: data = change(data)
            dest.writestr(name, data)
        for name, data in (extra or {}).items(): dest.writestr(name, data)
    return out.getvalue()


def test_interleaved_body_rows_cells_headers_paths_and_hashes():
    result = parse(fixture(), 10000)
    assert [x['locator'] for x in result] == ['Paragraph 1', 'Table 1, row 1', 'Table 1, row 2', 'Paragraph 2']
    assert result[2]['text'] == '2025 | (1,250.00)'
    assert result[1]['docx']['declared_header_row'] is True
    assert result[2]['docx']['cells'][1]['source_xml_path'] == '/w:document[1]/w:body[1]/w:tbl[1]/w:tr[2]/w:tc[2]'
    assert result[-1]['docx']['source_xml_path'].endswith('/w:p[2]')
    for item in result:
        m = item['docx']
        assert m['block_text_sha256'] == digest(item['text']) and m['parser_version'] == VERSION
        assert m['revision_handling'] == 'no_tracked_revisions_detected'
        assert not m['complete_document_verified']
    assert result == parse(fixture(), 10000)


def test_chunk_ranges_preserve_original_paragraph():
    doc = WordDocument(); doc.add_paragraph('Original ' + 'longer test ' * 700)
    out = io.BytesIO(); doc.save(out)
    result = parse_bytes(out.getvalue(), '.docx', 20000)
    text = ''.join(x['text'] for x in result)
    assert len(result) > 1
    for item in result:
        a, b = item['docx']['character_range']
        assert item['text'] == text[a:b] and item['docx']['block_text_sha256'] == digest(text)


@pytest.mark.parametrize('tag,code', [
    ('ins', 'tracked_revisions_require_resolution'), ('del', 'tracked_revisions_require_resolution'),
    ('tblGridChange', 'tracked_revisions_require_resolution'),
    ('moveFrom', 'tracked_revisions_require_resolution'), ('rPrChange', 'tracked_revisions_require_resolution'),
    ('commentRangeStart', 'comments_require_review'), ('fldSimple', 'fields_require_review'),
    ('drawing', 'unsupported_visual_or_embedded_content'), ('altChunk', 'unsupported_visual_or_embedded_content'),
    ('sdt', 'unsupported_body_structure'),
])
def test_unsupported_or_unresolved_body_is_not_silently_omitted(tag, code):
    raw = mutate(fixture(), lambda b: b.replace(b'<w:body>', b'<w:body><w:'+tag.encode()+b'/>', 1))
    with pytest.raises(DocxError) as error: parse(raw, 10000)
    assert error.value.code == code and error.value.path.startswith('/w:document[1]/w:body[1]')


def test_merged_cells_and_hidden_text_block():
    doc = WordDocument(); doc.add_paragraph('Original body.'); t = doc.add_table(rows=1, cols=2)
    t.cell(0,0).merge(t.cell(0,1)).text = 'Merged original text'
    out=io.BytesIO(); doc.save(out)
    with pytest.raises(DocxError, match='merged_cells_require_review'): parse(out.getvalue(),10000)
    raw = mutate(fixture(), lambda b: b.replace(b'<w:r>', b'<w:r><w:rPr><w:vanish/></w:rPr>',1))
    with pytest.raises(DocxError, match='hidden_text_requires_review'): parse(raw,10000)


def test_headers_and_comments_cannot_disappear():
    doc = WordDocument(); doc.add_paragraph('Original body.'); doc.sections[0].header.paragraphs[0].text='Original important header.'
    out=io.BytesIO(); doc.save(out)
    with pytest.raises(DocxError, match='supplemental_story_requires_review'): parse(out.getvalue(),10000)
    comment = ('<w:comments xmlns:w="'+W+'"><w:comment><w:p><w:r><w:t>Original comment</w:t></w:r></w:p></w:comment></w:comments>').encode()
    raw=mutate(fixture(),lambda b:b,extra={'word/comments.xml':comment})
    with pytest.raises(DocxError, match='comments_require_review'): parse(raw,10000)


def test_hyperlink_and_declared_numbering_preserved_without_fetching_or_inventing_number():
    doc=WordDocument(); p=doc.add_paragraph('Original linked paragraph: ')
    rel=doc.part.relate_to('https://example.invalid/no-fetch', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),rel)
    run=OxmlElement('w:r'); text=OxmlElement('w:t'); text.text='Visible original link'; run.append(text); link.append(run); p._p.append(link)
    num=p._p.get_or_add_pPr().get_or_add_numPr(); num.get_or_add_numId().val=7; num.get_or_add_ilvl().val=0
    out=io.BytesIO();doc.save(out)
    item=parse(out.getvalue(),10000)[0]
    assert 'Visible original link' in item['text']
    assert item['docx']['hyperlinks'][0]['target']=='https://example.invalid/no-fetch'
    assert item['docx']['declared_numbering']=={'ilvl':'0','numId':'7'}
    assert item['docx']['numbering_rendered'] is False


@pytest.mark.parametrize('kind', ['doctype','external','duplicate','traversal'])
def test_unsafe_packages(kind):
    raw=fixture()
    if kind=='doctype':
        raw=mutate(raw,lambda b:b.replace(b'<w:document',b'<!DOCTYPE x [<!ENTITY a "boom">]><w:document',1)); code='unsafe_xml'
    elif kind=='external':
        rel=b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="x" Type="attachedTemplate" TargetMode="External" Target="https://example.invalid/no-fetch"/></Relationships>'
        raw=mutate(raw,lambda b:rel,part='word/_rels/document.xml.rels');code='active_or_embedded_content'
    elif kind=='duplicate':
        with pytest.warns(UserWarning,match='Duplicate name'):
            raw=mutate(raw,lambda b:b,extra={'word/document.xml':b'bad'})
        code='archive_limit_or_ambiguous_members'
    else:
        raw=mutate(raw,lambda b:b,extra={'../escape.xml':b'<x/>'});code='unsafe_archive_path'
    with pytest.raises(DocxError,match=code):parse(raw,10000)


def test_empty_cells_and_paragraph_boundaries_are_preserved():
    doc=WordDocument(); t=doc.add_table(rows=1,cols=2); t.cell(0,0).text='First | value';t.cell(0,0).add_paragraph('Second paragraph')
    out=io.BytesIO();doc.save(out)
    item=parse(out.getvalue(),10000)[0]
    assert item['text']=='First \\| value\\nSecond paragraph | '
    assert item['docx']['cells'][1]['text']==''
    assert len(item['docx']['cells'][0]['paragraphs'])==2


def test_upload_failure_is_atomic_and_success_retains_order_and_limits(client):
    activate(client)
    raw=mutate(fixture(),lambda b:b.replace(b'<w:body>',b'<w:body><w:ins/>',1))
    bad=client.post('/api/v1/workspaces/demo-workspace/documents', files={'file':('revision.docx',raw,MIME)},data={'authorization_basis':'own_original'})
    assert bad.status_code==422, bad.text
    assert bad.json()['error']['code']=='DOCX_PARSE_BLOCKED'
    with client.app.state.db.Session() as db: assert db.scalar(select(func.count()).select_from(Document))==0
    good=client.post('/api/v1/workspaces/demo-workspace/documents',files={'file':('original.docx',fixture(),MIME)},data={'authorization_basis':'own_original'})
    assert good.status_code==201,good.text
    did=good.json()['id']; detail=client.get('/api/v1/documents/'+did).json()
    assert detail['chunks'][1]['locator']=='Table 1, row 1'
    run=SimpleNamespace(document_ids=[did],workspace_id='demo-workspace',inputs={},workflow='document_gaap',context={})
    with client.app.state.db.Session() as db:
        limits=preprocess(db,run,{'min_documents':1,'guardrail':'Synthetic'})['document_extraction_limits'][0]
        assert limits['parser_versions']==[VERSION] and limits['body_order_metadata']=='available'
        doc=db.get(Document,did);doc.chunks=[{'locator':'Paragraph 1','text':'Original legacy text.'}];db.flush()
        legacy=preprocess(db,run,{'min_documents':1,'guardrail':'Synthetic'})['document_extraction_limits'][0]
        assert legacy['body_order_metadata']=='unknown_legacy_extraction'


def test_empty_table_does_not_count_separators_as_source_text():
    doc=WordDocument();doc.add_table(rows=2,cols=10)
    out=io.BytesIO();doc.save(out)
    with pytest.raises(DocxError,match='no_text'):parse(out.getvalue(),10000)


def test_missing_package_part_is_not_treated_as_complete_body():
    raw=mutate(fixture(),lambda b:b.replace(b'Target="styles.xml"',b'Target="missing-styles.xml"'),part='word/_rels/document.xml.rels')
    with pytest.raises(DocxError,match='missing_relationship_part'):parse(raw,10000)


def test_duplicate_relationship_identity_is_rejected():
    import xml.etree.ElementTree as ET
    def duplicate(raw):
        root=ET.fromstring(raw)
        root.append(ET.fromstring(ET.tostring(root[0])))
        return ET.tostring(root)
    raw=mutate(fixture(),duplicate,part='word/_rels/document.xml.rels')
    with pytest.raises(DocxError,match='ambiguous_relationship'):parse(raw,10000)
